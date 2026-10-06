"""Data acquisition, validation, cleaning, and weekly panel aggregation."""

from __future__ import annotations

import datetime
import hashlib
import json
import urllib.request
import zipfile
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import yaml

KNOWN_NONMERCHANDISE_CODES = {
    "POST",
    "D",
    "M",
    "BANK CHARGES",
    "PADS",
    "DOT",
    "CRUK",
    "AMAZONFEE",
    "TEST",
    "ADJUST",
    "ADJUST2",
    "S",
    "B",
    "GIFT",
}

HEADER_ALIASES = {
    "invoice": "invoice_id",
    "invoiceno": "invoice_id",
    "stockcode": "sku_id",
    "description": "description",
    "quantity": "quantity",
    "invoicedate": "transaction_at",
    "price": "sale_unit_price_gbp",
    "unitprice": "sale_unit_price_gbp",
    "customer id": "customer_id",
    "customerid": "customer_id",
    "country": "country",
}


def compute_sha256(file_path: str | Path) -> str:
    """Compute SHA-256 hex digest of a file."""
    hasher = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def acquire_source_data(config_path: str = "config/project.yaml") -> dict[str, Any]:
    """Download source archive if needed, extract, inspect sheets, and record manifest."""
    with open(config_path, "r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    data_cfg = cfg["data"]
    raw_excel_path = Path(data_cfg["raw_excel_path"])
    raw_zip_path = Path(data_cfg.get("raw_zip_path", "data/raw/online_retail_II.zip"))
    manifest_path = Path(data_cfg["source_manifest_path"])

    raw_excel_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Download if excel doesn't exist
    if not raw_excel_path.exists():
        if not raw_zip_path.exists():
            download_url = data_cfg["source_download_url"]
            print(f"Downloading source dataset from {download_url}...")
            # Use headers to prevent 403 on some mirrors
            req = urllib.request.Request(
                download_url,
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) DemandGuard/0.1"},
            )
            with (
                urllib.request.urlopen(req, timeout=120) as resp,
                open(raw_zip_path, "wb") as out_f,
            ):
                out_f.write(resp.read())
            print(f"Downloaded archive to {raw_zip_path} ({raw_zip_path.stat().st_size} bytes)")

        # Unzip
        print(f"Extracting {raw_zip_path}...")
        with zipfile.ZipFile(raw_zip_path, "r") as z:
            for item in z.namelist():
                if item.endswith(".xlsx"):
                    with z.open(item) as zf, open(raw_excel_path, "wb") as out_f:
                        out_f.write(zf.read())
                    break
        print(f"Extracted excel to {raw_excel_path} ({raw_excel_path.stat().st_size} bytes)")

    # 2. Inspect Excel workbook structure
    wb = openpyxl.load_workbook(raw_excel_path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    sheet_stats = {}
    total_raw_rows = 0
    column_maps = {}

    for sheet in sheet_names:
        df_sample = pd.read_excel(raw_excel_path, sheet_name=sheet, nrows=5)
        raw_cols = list(df_sample.columns)
        norm_cols = {c: HEADER_ALIASES.get(str(c).strip().lower(), str(c)) for c in raw_cols}
        column_maps[sheet] = norm_cols

        # Count rows in sheet
        # Fast count using openpyxl
        wb_s = openpyxl.load_workbook(raw_excel_path, read_only=True)
        ws = wb_s[sheet]
        row_count = ws.max_row - 1 if ws.max_row else 0
        wb_s.close()
        sheet_stats[sheet] = {"row_count": row_count, "columns": raw_cols}
        total_raw_rows += row_count

    file_size = raw_excel_path.stat().st_size
    sha256 = compute_sha256(raw_excel_path)

    manifest = {
        "dataset_name": "Online Retail II",
        "source_landing_url": data_cfg["source_landing_url"],
        "source_download_url": data_cfg["source_download_url"],
        "retrieved_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "sha256": sha256,
        "byte_size": file_size,
        "license": "CC BY 4.0",
        "sheet_names": sheet_names,
        "sheet_stats": sheet_stats,
        "total_raw_rows": total_raw_rows,
        "column_mappings": column_maps,
    }

    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)

    print(f"Source manifest written to {manifest_path} (SHA-256: {sha256[:16]}...)")
    return manifest


def load_raw_sheets_to_df(raw_excel_path: Path | str) -> pd.DataFrame:
    """Load all sheets and standardise column names."""
    excel_file = pd.ExcelFile(raw_excel_path)
    dfs = []
    for sheet in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet)
        col_map = {c: HEADER_ALIASES.get(str(c).strip().lower(), str(c)) for c in df.columns}
        df = df.rename(columns=col_map)
        df["_source_sheet"] = sheet
        df["_source_row"] = df.index + 2
        dfs.append(df)

    combined = pd.concat(dfs, ignore_index=True)
    return combined


def clean_transactions(
    df_raw: pd.DataFrame,
    country: str = "United Kingdom",
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Apply cleaning rules and return clean dataframe plus exclusion audit."""
    audit = {
        "initial_raw_rows": len(df_raw),
        "initial_raw_units": int(df_raw["quantity"].sum()) if "quantity" in df_raw else 0,
        "exclusions": {},
    }

    df = df_raw.copy()

    # 1. Missing essential identifiers/dates
    missing_mask = (
        df["invoice_id"].isna()
        | df["sku_id"].isna()
        | df["transaction_at"].isna()
        | df["quantity"].isna()
        | df["sale_unit_price_gbp"].isna()
    )
    audit["exclusions"]["missing_essentials"] = {
        "rows": int(missing_mask.sum()),
        "units": int(df.loc[missing_mask, "quantity"].fillna(0).sum()),
    }
    df = df[~missing_mask].copy()

    # Cast types
    df["invoice_id"] = df["invoice_id"].astype(str).str.strip()
    df["sku_id"] = df["sku_id"].astype(str).str.strip()
    df["quantity"] = pd.to_numeric(df["quantity"], errors="coerce").fillna(0).astype(int)
    df["sale_unit_price_gbp"] = pd.to_numeric(df["sale_unit_price_gbp"], errors="coerce")
    df["transaction_at"] = pd.to_datetime(df["transaction_at"])
    df["country"] = df["country"].astype(str).str.strip()

    # 2. Cancellations (Invoice starts with C/c)
    canc_mask = df["invoice_id"].str.upper().str.startswith("C")
    audit["exclusions"]["cancellations"] = {
        "rows": int(canc_mask.sum()),
        "units": int(df.loc[canc_mask, "quantity"].sum()),
    }
    df = df[~canc_mask].copy()

    # 3. Nonpositive quantities (<= 0)
    nonpos_qty_mask = df["quantity"] <= 0
    audit["exclusions"]["nonpositive_quantity"] = {
        "rows": int(nonpos_qty_mask.sum()),
        "units": int(df.loc[nonpos_qty_mask, "quantity"].sum()),
    }
    df = df[~nonpos_qty_mask].copy()

    # 4. Nonpositive / nonfinite prices (<= 0)
    nonpos_price_mask = (df["sale_unit_price_gbp"] <= 0) | df["sale_unit_price_gbp"].isna()
    audit["exclusions"]["nonpositive_or_nan_price"] = {
        "rows": int(nonpos_price_mask.sum()),
        "units": int(df.loc[nonpos_price_mask, "quantity"].sum()),
    }
    df = df[~nonpos_price_mask].copy()

    # 5. Nonmerchandise service codes
    nonmerch_mask = df["sku_id"].str.upper().isin(KNOWN_NONMERCHANDISE_CODES)
    audit["exclusions"]["nonmerchandise_codes"] = {
        "rows": int(nonmerch_mask.sum()),
        "units": int(df.loc[nonmerch_mask, "quantity"].sum()),
    }
    df = df[~nonmerch_mask].copy()

    # 6. Country filter
    country_mask = df["country"].str.lower() != country.lower()
    audit["exclusions"]["non_target_country"] = {
        "rows": int(country_mask.sum()),
        "units": int(df.loc[country_mask, "quantity"].sum()),
    }
    df = df[~country_mask].copy()

    # 7. Exact duplicate rows audit
    dup_cols = ["invoice_id", "sku_id", "quantity", "transaction_at", "sale_unit_price_gbp"]
    dup_mask = df.duplicated(subset=dup_cols, keep="first")
    audit["exclusions"]["exact_duplicates_retained_as_legit_transactions"] = {
        "rows": int(dup_mask.sum()),
        "units": int(df.loc[dup_mask, "quantity"].sum()),
    }

    audit["final_clean_rows"] = len(df)
    audit["final_clean_units"] = int(df["quantity"].sum())

    # Keep only canonical columns
    clean_df = df[
        [
            "invoice_id",
            "sku_id",
            "description",
            "quantity",
            "transaction_at",
            "sale_unit_price_gbp",
            "country",
        ]
    ].copy()

    return clean_df, audit


def prepare_weekly_panel(
    df_clean: pd.DataFrame,
    output_parquet_path: Path | str = "data/processed/weekly_sales.parquet",
) -> pd.DataFrame:
    """Aggregate cleaned transactions to weekly product sales using Monday-start complete weeks."""
    con = duckdb.connect(database=":memory:")
    con.register("clean_tx", df_clean)

    # 1. Compute week boundaries and identify first/last partial weeks
    # Monday is day of week 1 in ISO.
    query_dates = """
        SELECT
            min(transaction_at) as min_tx,
            max(transaction_at) as max_tx
        FROM clean_tx
    """
    date_res = con.execute(query_dates).df()
    min_tx = pd.to_datetime(date_res["min_tx"].iloc[0])
    max_tx = pd.to_datetime(date_res["max_tx"].iloc[0])

    # First Monday at or after min_tx
    first_monday = min_tx.floor("D")
    if first_monday.weekday() != 0:  # 0 is Monday
        first_monday = first_monday + pd.Timedelta(days=(7 - first_monday.weekday()))

    # Last Sunday at or before max_tx
    last_sunday = max_tx.floor("D")
    if last_sunday.weekday() != 6:  # 6 is Sunday
        last_sunday = last_sunday - pd.Timedelta(days=(last_sunday.weekday() + 1))
    last_monday = last_sunday - pd.Timedelta(days=6)

    print(f"Source range: {min_tx} to {max_tx}")
    print(
        f"Complete week range: {first_monday.date()} to {last_sunday.date()} (Last Monday: {last_monday.date()})"
    )

    # 2. DuckDB aggregation to Monday week_start
    query_weekly = f"""
        WITH tx_weeks AS (
            SELECT
                sku_id,
                -- Truncate to Monday
                CAST(date_trunc('week', transaction_at) AS DATE) AS week_start,
                quantity
            FROM clean_tx
            WHERE transaction_at >= TIMESTAMP '{first_monday.strftime("%Y-%m-%d %H:%M:%S")}'
              AND transaction_at < TIMESTAMP '{(last_sunday + pd.Timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")}'
        )
        SELECT
            sku_id,
            week_start,
            SUM(quantity) AS units_sold
        FROM tx_weeks
        GROUP BY sku_id, week_start
        ORDER BY sku_id, week_start
    """
    weekly_df = con.execute(query_weekly).df()
    con.close()

    # Save to Parquet
    out_path = Path(output_parquet_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    table = pa.Table.from_pandas(weekly_df)
    pq.write_table(table, out_path)
    print(
        f"Saved weekly sales panel to {out_path} ({len(weekly_df)} rows, {weekly_df['sku_id'].nunique()} unique SKUs)"
    )
    return weekly_df


def run_prepare_pipeline(config_path: str = "config/project.yaml") -> dict[str, Any]:
    """Execute complete prepare workflow from raw Excel to cleaned parquet and quality report."""
    with open(config_path, "r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    data_cfg = cfg["data"]
    raw_excel_path = Path(data_cfg["raw_excel_path"])
    weekly_sales_path = Path(data_cfg["weekly_sales_path"])

    print("Loading raw Excel sheets...")
    df_raw = load_raw_sheets_to_df(raw_excel_path)
    print(f"Loaded {len(df_raw)} raw transaction rows across all sheets.")

    print("Cleaning transactions according to data specification...")
    df_clean, audit = clean_transactions(df_raw, country=data_cfg.get("country", "United Kingdom"))

    print("Aggregating complete weekly sales panel...")
    weekly_df = prepare_weekly_panel(df_clean, output_parquet_path=weekly_sales_path)

    # Write reports/data_quality.md and audit json
    rep_dir = Path("reports")
    rep_dir.mkdir(parents=True, exist_ok=True)

    with open(rep_dir / "data_quality_audit.json", "w", encoding="utf-8") as f:
        json.dump(audit, f, indent=2)

    total_panel_units = int(weekly_df["units_sold"].sum())

    report_md = f"""# Data Quality and Cleaning Report

## 1. Raw Source Summary
- Raw Rows: {audit["initial_raw_rows"]:,}
- Raw Units: {audit["initial_raw_units"]:,}

## 2. Exclusions Breakdown
| Exclusion Reason | Rows Excluded | Units Excluded |
| --- | --- | --- |
| Missing Identifiers/Dates | {audit["exclusions"]["missing_essentials"]["rows"]:,} | {audit["exclusions"]["missing_essentials"]["units"]:,} |
| Cancellations | {audit["exclusions"]["cancellations"]["rows"]:,} | {audit["exclusions"]["cancellations"]["units"]:,} |
| Nonpositive Quantity (<=0) | {audit["exclusions"]["nonpositive_quantity"]["rows"]:,} | {audit["exclusions"]["nonpositive_quantity"]["units"]:,} |
| Nonpositive/NaN Price | {audit["exclusions"]["nonpositive_or_nan_price"]["rows"]:,} | {audit["exclusions"]["nonpositive_or_nan_price"]["units"]:,} |
| Nonmerchandise Codes | {audit["exclusions"]["nonmerchandise_codes"]["rows"]:,} | {audit["exclusions"]["nonmerchandise_codes"]["units"]:,} |
| Non-UK Transactions | {audit["exclusions"]["non_target_country"]["rows"]:,} | {audit["exclusions"]["non_target_country"]["units"]:,} |

## 3. Clean Transactions and Aggregation
- Final Clean Transaction Rows: {audit["final_clean_rows"]:,}
- Final Clean Units: {audit["final_clean_units"]:,}
- Prepared Panel Nonzero Rows: {len(weekly_df):,}
- Prepared Panel Total Units: {total_panel_units:,}
- Unique Clean SKUs: {weekly_df["sku_id"].nunique():,}
"""
    with open(rep_dir / "data_quality.md", "w", encoding="utf-8") as f:
        f.write(report_md)

    print(f"Data quality report saved to {rep_dir / 'data_quality.md'}")
    return {"audit": audit, "panel_rows": len(weekly_df), "panel_units": total_panel_units}
